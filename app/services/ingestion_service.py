import io
import hashlib
from pathlib import Path
from typing import Optional, Tuple
import pymupdf
from PIL import Image

from app.settings import settings
from app.domain.errors import (
    DocumentoInvalidoError,
    TamanoArchivoExcedidoError,
    DocumentoCorruptoOEncriptadoError,
    ExcesoPaginasError,
)

PDF_MAGIC_BYTES = b"%PDF"
SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".tiff",
    ".tif",
    ".docx",
    ".xlsx",
    ".txt",
}


def _convert_docx_to_pdf_in_memory(docx_bytes: bytes) -> bytes:
    """Extrae texto y tablas de un archivo DOCX y genera un PDF vectorial en memoria con PyMuPDF."""
    import docx
    doc_in = docx.Document(io.BytesIO(docx_bytes))
    pdf_doc = pymupdf.open()
    page = pdf_doc.new_page(width=595, height=842)  # A4

    y_cursor = 50.0
    margin_x = 50.0
    line_height = 14.0

    def check_new_page(needed_height=20.0):
        nonlocal page, y_cursor
        if y_cursor + needed_height > 790.0:
            page = pdf_doc.new_page(width=595, height=842)
            y_cursor = 50.0

    for p in doc_in.paragraphs:
        txt = p.text.strip()
        if not txt:
            y_cursor += line_height * 0.5
            continue
        check_new_page(line_height + 4)
        is_heading = p.style.name.startswith("Heading") if p.style else False
        fsize = 13.0 if is_heading else 9.5
        page.insert_text(pymupdf.Point(margin_x, y_cursor), txt[:140], fontsize=fsize, fontname="helv")
        y_cursor += line_height + (4 if is_heading else 1)

    for table in doc_in.tables:
        check_new_page(30.0)
        for row in table.rows:
            check_new_page(line_height + 4)
            row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_txt:
                page.insert_text(pymupdf.Point(margin_x + 10, y_cursor), row_txt[:120], fontsize=9.0, fontname="helv")
                y_cursor += line_height + 2

    res_bytes = pdf_doc.tobytes()
    pdf_doc.close()
    return res_bytes


def _convert_xlsx_to_pdf_in_memory(xlsx_bytes: bytes) -> bytes:
    """Extrae datos de hojas de cálculo XLSX y genera un PDF estructurado en memoria."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    pdf_doc = pymupdf.open()
    page = pdf_doc.new_page(width=842, height=595)  # A4 Landscape para tablas

    margin_x = 40.0
    y_cursor = 50.0
    line_height = 14.0

    def check_new_page(needed_height=20.0):
        nonlocal page, y_cursor
        if y_cursor + needed_height > 550.0:
            page = pdf_doc.new_page(width=842, height=595)
            y_cursor = 50.0

    for sheet_name in wb.sheetnames[:3]:
        sheet = wb[sheet_name]
        check_new_page(25.0)
        page.insert_text(pymupdf.Point(margin_x, y_cursor), f"HOJA: {sheet_name.upper()}", fontsize=11.0, fontname="helv")
        y_cursor += line_height + 4

        for row in sheet.iter_rows(values_only=True):
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals:
                check_new_page(line_height + 2)
                if len(vals) == 2:
                    k = vals[0] if vals[0].endswith(":") else f"{vals[0]}:"
                    row_str = f"{k} {vals[1]}"
                else:
                    row_str = "    ".join(vals)
                page.insert_text(pymupdf.Point(margin_x, y_cursor), row_str[:160], fontsize=9.0, fontname="helv")
                y_cursor += line_height

    res_bytes = pdf_doc.tobytes()
    pdf_doc.close()
    return res_bytes


def _convert_txt_to_pdf_in_memory(txt_bytes: bytes) -> bytes:
    """Convierte texto plano a un documento PDF vectorial paginado en memoria."""
    text_content = txt_bytes.decode("utf-8", errors="replace")
    lines = text_content.splitlines()

    pdf_doc = pymupdf.open()
    page = pdf_doc.new_page(width=595, height=842)
    margin_x = 50.0
    y_cursor = 50.0
    line_height = 13.0

    for line in lines:
        if y_cursor + line_height > 790.0:
            page = pdf_doc.new_page(width=595, height=842)
            y_cursor = 50.0
        page.insert_text(pymupdf.Point(margin_x, y_cursor), line[:130], fontsize=9.5, fontname="helv")
        y_cursor += line_height

    res_bytes = pdf_doc.tobytes()
    pdf_doc.close()
    return res_bytes


def _convert_image_to_pdf_in_memory(img_bytes: bytes, ext: str) -> bytes:
    """Convierte una imagen raster (PNG, JPG, TIFF, WEBP) a una página de documento PDF en memoria nativa en C."""
    clean_ext = ext.lstrip(".").lower()
    if clean_ext == "tif":
        clean_ext = "tiff"
    elif clean_ext == "jpg":
        clean_ext = "jpeg"

    try:
        img_doc = pymupdf.open(stream=img_bytes, filetype=clean_ext)
        pdf_bytes = img_doc.convert_to_pdf()
        img_doc.close()
        return pdf_bytes
    except Exception:
        # Fallback a través de Pillow si el formato requiere descomposición
        pil_img = Image.open(io.BytesIO(img_bytes))
        if pil_img.mode not in ("RGB", "L"):
            pil_img = pil_img.convert("RGB")
        buf = io.BytesIO()
        pil_img.save(buf, format="PNG")
        png_bytes = buf.getvalue()
        img_doc = pymupdf.open(stream=png_bytes, filetype="png")
        pdf_bytes = img_doc.convert_to_pdf()
        img_doc.close()
        return pdf_bytes


def validate_and_read_pdf(
    pdf_bytes: bytes,
    filename: Optional[str] = None,
    max_file_size: Optional[int] = None,
    max_pages: Optional[int] = None,
) -> Tuple[str, pymupdf.Document, int]:
    """
    Valida e ingesta un documento en memoria, soportando de forma universal:
    - PDFs nativos (.pdf)
    - Imágenes forenses (.png, .jpg, .jpeg, .webp, .tiff, .tif)
    - Documentos estructurados (.docx, .xlsx, .txt)

    Virtualiza cualquier formato en un documento pymupdf.Document de alta velocidad
    para reutilizar el 100% de la arquitectura existente (clasificación, OCR selectivo,
    bounding boxes, visor interactivo y chat).

    Retorna:
        Tuple[str, pymupdf.Document, int]: (doc_hash, fitz_doc, total_paginas)
    """
    limit_bytes = max_file_size or settings.MAX_FILE_SIZE_BYTES
    limit_pages = max_pages or settings.MAX_PAGES_PER_DOCUMENT

    # 1. Validación de tamaño y contenido no vacío
    if not pdf_bytes or len(pdf_bytes) == 0:
        raise DocumentoInvalidoError("El archivo provisto está vacío.")

    received_bytes = len(pdf_bytes)
    if received_bytes > limit_bytes:
        raise TamanoArchivoExcedidoError(max_bytes=limit_bytes, received_bytes=received_bytes)

    # 2. Detección de extensión
    ext = ".pdf"
    if filename:
        ext = Path(filename).suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
            raise DocumentoInvalidoError(f"Formato '{ext}' no soportado. Formatos admitidos: {allowed}")

    # 3. Cálculo determinista de hash SHA-256 sobre los bytes originales
    doc_hash = hashlib.sha256(pdf_bytes).hexdigest()

    # 4. Normalización a Documento PyMuPDF en memoria según el tipo de archivo
    working_pdf_bytes = pdf_bytes

    try:
        if ext in (".png", ".jpg", ".jpeg", ".webp", ".tiff", ".tif"):
            working_pdf_bytes = _convert_image_to_pdf_in_memory(pdf_bytes, ext)
        elif ext == ".docx":
            working_pdf_bytes = _convert_docx_to_pdf_in_memory(pdf_bytes)
        elif ext == ".xlsx":
            working_pdf_bytes = _convert_xlsx_to_pdf_in_memory(pdf_bytes)
        elif ext == ".txt":
            working_pdf_bytes = _convert_txt_to_pdf_in_memory(pdf_bytes)
        else:
            # Para PDFs, validar cabecera %PDF
            header = pdf_bytes[:1024]
            if PDF_MAGIC_BYTES not in header:
                raise DocumentoInvalidoError("El archivo provisto no es un PDF válido o no cuenta con la cabecera %PDF.")

        doc = pymupdf.open(stream=working_pdf_bytes, filetype="pdf")
    except DocumentoInvalidoError:
        raise
    except Exception as exc:
        raise DocumentoCorruptoOEncriptadoError(
            f"El documento está corrupto o mal formado y no pudo ser procesado: {str(exc)}"
        ) from exc

    # 5. Verificación de cifrado / contraseña
    if doc.is_encrypted or doc.needs_pass:
        doc.close()
        raise DocumentoCorruptoOEncriptadoError(
            "El documento está protegido con contraseña y no puede ser procesado sin credenciales."
        )

    # 6. Verificación de conteo de páginas
    page_count = doc.page_count
    if page_count == 0:
        doc.close()
        raise DocumentoInvalidoError("El documento no contiene páginas válidas.")

    if page_count > limit_pages:
        doc.close()
        raise ExcesoPaginasError(max_pages=limit_pages, total_pages=page_count)

    return doc_hash, doc, page_count
